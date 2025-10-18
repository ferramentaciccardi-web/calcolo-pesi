# app.py
import streamlit as st
import pandas as pd
from io import StringIO
from openpyxl import load_workbook

# Scanner (camera)
from streamlit_webrtc import webrtc_streamer, VideoTransformerBase
import av
import numpy as np

st.set_page_config(page_title="Calcolatore Peso Prodotti", layout="wide")
st.title("⚖️ Calcolatore Peso fuochi")

# ---------------- FUNZIONI ----------------
def read_products_from_excel(file):
    """
    Legge il file Excel e ritorna un DataFrame con:
    Codice (col A), Nome (col B), Flag (col C), Peso_unitario_kg (col I),
    CodiceEAN (col J se presente; altrimenti vuoto).
    Dati a partire dalla riga 3.
    """
    wb = load_workbook(file, data_only=True, read_only=True)
    if "Foglio 2" not in wb.sheetnames:
        st.error(f"'Foglio 2' non trovato. Fogli disponibili: {wb.sheetnames}")
        return pd.DataFrame()
    ws = wb["Foglio 2"]

    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        codice = row[0] if len(row) > 0 else None
        nome = row[1] if len(row) > 1 else None
        flag = row[2] if len(row) > 2 else None
        peso_unitario = row[8] if len(row) > 8 else None  # colonna I
        codice_ean = row[9] if len(row) > 9 else None     # colonna J (opzionale)

        # salta righe completamente vuote
        if all(x in (None, "") for x in [codice, nome, peso_unitario, codice_ean]):
            continue

        rows.append({
            "Codice": str(codice).strip() if codice not in (None, "") else "",
            "Nome": str(nome).strip() if nome not in (None, "") else "",
            "Flag": str(flag).strip() if flag not in (None, "") else "",
            "Peso_unitario_kg": float(peso_unitario) if str(peso_unitario).strip() not in ("", "None") and peso_unitario is not None else None,
            "CodiceEAN": str(codice_ean).strip() if codice_ean not in (None, "") else ""
        })

    df = pd.DataFrame(rows)
    if "CodiceEAN" not in df.columns:
        df["CodiceEAN"] = ""
    return df


def calcola_peso(unit_weight, pezzi, flag):
    try:
        peso = float(unit_weight) * float(pezzi)
        if "√" in str(flag):
            peso /= 2
        return peso
    except Exception:
        return 0.0


# ---------------- STATO ----------------
ss = st.session_state
ss.setdefault("carrello", [])
ss.setdefault("prodotti_df", pd.DataFrame())
ss.setdefault("lista_prodotti", [])
ss.setdefault("selected_index", None)
ss.setdefault("scanned_code", "")

# ---------------- UPLOAD EXCEL ----------------
uploaded_file = st.file_uploader("📁 Carica file Excel (richiesto il foglio 'Foglio 2')", type=["xlsx", "xlsm"])
if uploaded_file:
    ss.prodotti_df = read_products_from_excel(uploaded_file)
    if ss.prodotti_df.empty:
        st.warning("Il foglio 'Foglio 2' non contiene prodotti validi.")
    else:
        st.success(f"✅ File caricato: trovati {len(ss.prodotti_df)} prodotti.")
        ss.lista_prodotti = [f"{row.Codice} - {row.Nome}" for _, row in ss.prodotti_df.iterrows()]

st.divider()
st.subheader("📷 Scanner codici (usa colonna 'CodiceEAN')")

scanner_col1, scanner_col2 = st.columns([2, 1])

with scanner_col1:
    st.caption("Inquadra codice a barre/QR. Se riconosciuto, compila il campo qui a destra e seleziona il prodotto EAN corrispondente.")
    # Video transformer con pyzbar (se disponibile)
    class BarcodeDetector(VideoTransformerBase):
        def _init_(self):
            self.last_text = None

        def transform(self, frame: av.VideoFrame):
            image = frame.to_ndarray(format="bgr24")
            try:
                from pyzbar.pyzbar import decode
                decoded = decode(image)
                if decoded:
                    text = decoded[0].data.decode("utf-8").strip()
                    self.last_text = text
                    ss.scanned_code = text
            except Exception:
                # Se pyzbar/zbar non disponibili in questo ambiente, non blocchiamo la UI.
                pass
            return image

    webrtc_streamer(
        key="barcode-scanner",
        video_transformer_factory=BarcodeDetector,
        media_stream_constraints={"video": True, "audio": False},
        async_processing=True,
    )

with scanner_col2:
    st.text_input("Ultimo codice rilevato (EAN/QR)", key="scanned_code", disabled=True)
    if ss.scanned_code and not ss.prodotti_df.empty:
        df = ss.prodotti_df
        matches = df.index[df["CodiceEAN"] == ss.scanned_code].tolist()
        if matches:
            st.success("✅ Codice EAN trovato.")
            ss.selected_index = matches[0]
        else:
            st.info("⚠️ EAN non trovato nel file caricato.")

# ---------------- SELEZIONE PRODOTTO & INPUT ----------------
st.divider()
if not ss.prodotti_df.empty:
    left, right = st.columns([1, 2])

    with left:
        options = [""] + ss.lista_prodotti
        # selezione sincronizzata con scanner
        index_to_show = 0
        if ss.selected_index is not None:
            index_to_show = ss.selected_index + 1  # +1 per l'opzione vuota
        selected = st.selectbox("Seleziona prodotto", options, index=index_to_show)

    # determina "prodotto" corrente
    prodotto = None
    if selected and index_to_show != 0:
        prodotto = ss.prodotti_df.iloc[index_to_show - 1]
    elif selected:
        idx = options.index(selected) - 1
        prodotto = ss.prodotti_df.iloc[idx]
        ss.selected_index = idx
    elif ss.selected_index is not None:
        prodotto = ss.prodotti_df.iloc[ss.selected_index]

    if prodotto is not None:
        codice = prodotto["Codice"]
        nome = prodotto["Nome"]
        flag = prodotto["Flag"]
        peso_unit = float(prodotto["Peso_unitario_kg"] or 0)
        ean = prodotto.get("CodiceEAN", "")
        st.caption(f"Codice interno: {codice} | EAN: {ean}")
    else:
        # fallback inserimento manuale
        codice = st.text_input("Codice interno")
        nome = st.text_input("Nome prodotto")
        flag = st.text_input("Flag (usa '√' per dimezzare)")
        peso_unit = st.number_input("Peso unitario (kg)", min_value=0.0, step=0.01)

    pezzi = st.number_input("Numero pezzi", min_value=0.0, step=1.0)

    if st.button("🛒 Aggiungi al riepilogo"):
        peso_tot = calcola_peso(peso_unit, pezzi, flag)
        ss.carrello.append({
            "Codice": codice,
            "Nome": nome,
            "Flag": flag,
            "Peso_unitario_kg": peso_unit,
            "Pezzi": pezzi,
            "Kg": peso_tot
        })
        st.success("Prodotto aggiunto!")

# ---------------- RIEPILOGO ----------------
if ss.carrello:
    st.subheader("📋 Riepilogo prodotti")
    df_cart = pd.DataFrame(ss.carrello)
    df_cart.index += 1
    st.dataframe(df_cart, use_container_width=True)

    totale = df_cart["Kg"].sum()
    st.metric("Totale peso (kg)", f"{totale:.4f}")

    # Download CSV
    csv_buffer = StringIO()
    df_cart.to_csv(csv_buffer, index=False)
    st.download_button(
        label="💾 Scarica riepilogo CSV",
        data=csv_buffer.getvalue(),
        file_name="riepilogo_pesi.csv",
        mime="text/csv"
    )

    if st.button("🗑️ Svuota riepilogo"):
        ss.carrello = []
        ss.selected_index = None
        ss.scanned_code = ""
        st.rerun()

# ---------------- INFO ----------------
if ss.prodotti_df.empty:
    st.info("Carica un file Excel per iniziare. Assicurati di avere (se possibile) la colonna *J: CodiceEAN* con i codici a barre.")
